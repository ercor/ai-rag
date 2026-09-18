"""Unit tests for parser block contract.

Each parser must return ParsedDocument with appropriate block types:
- PDF: PageMarker blocks between pages, Paragraph blocks for text
- DOCX: Heading, Paragraph, ListBlock, TableBlock blocks
- HTML: Heading, Paragraph, ListBlock, TableBlock blocks
- TXT/MD: Heading (from # prefix), Paragraph blocks
"""

import pytest

from src.ports.models import (
    Heading,
    ListBlock,
    PageMarker,
    Paragraph,
    ParsedDocument,
    TableBlock,
)


class TestPDFParserBlocks:
    """PDF parser returns ParsedDocument with PageMarker blocks."""

    async def test_pdf_returns_parsed_document(self, tmp_path):
        """PDF parser returns ParsedDocument, not str."""
        from src.adapters.parsers import MultiFormatParser

        # Create a minimal PDF using pymupdf
        try:
            import pymupdf
        except ImportError:
            pytest.skip("pymupdf not installed")

        pdf_path = tmp_path / "test.pdf"
        doc = pymupdf.open()
        page1 = doc.new_page()
        page1.insert_text((72, 72), "First page content.")
        page2 = doc.new_page()
        page2.insert_text((72, 72), "Second page content.")
        doc.save(str(pdf_path))
        doc.close()

        parser = MultiFormatParser()
        result = await parser.parse(str(pdf_path))

        assert isinstance(result, ParsedDocument)
        assert len(result.blocks) > 0

    async def test_pdf_has_page_markers(self, tmp_path):
        """PDF parser produces PageMarker blocks for each page."""
        from src.adapters.parsers import MultiFormatParser

        try:
            import pymupdf
        except ImportError:
            pytest.skip("pymupdf not installed")

        pdf_path = tmp_path / "pages.pdf"
        doc = pymupdf.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {i + 1} text content.")
        doc.save(str(pdf_path))
        doc.close()

        parser = MultiFormatParser()
        result = await parser.parse(str(pdf_path))

        page_markers = [b for b in result.blocks if isinstance(b, PageMarker)]
        assert len(page_markers) == 3
        assert page_markers[0].page_number == 1
        assert page_markers[1].page_number == 2
        assert page_markers[2].page_number == 3

    async def test_pdf_has_paragraph_blocks(self, tmp_path):
        """PDF parser produces Paragraph blocks for text content."""
        from src.adapters.parsers import MultiFormatParser

        try:
            import pymupdf
        except ImportError:
            pytest.skip("pymupdf not installed")

        pdf_path = tmp_path / "text.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Hello world paragraph.")
        doc.save(str(pdf_path))
        doc.close()

        parser = MultiFormatParser()
        result = await parser.parse(str(pdf_path))

        paragraphs = [b for b in result.blocks if isinstance(b, Paragraph)]
        assert len(paragraphs) >= 1
        assert any("Hello world" in p.text for p in paragraphs)

    async def test_pdf_metadata_has_content_type(self, tmp_path):
        """PDF parser sets content_type in metadata."""
        from src.adapters.parsers import MultiFormatParser

        try:
            import pymupdf
        except ImportError:
            pytest.skip("pymupdf not installed")

        pdf_path = tmp_path / "meta.pdf"
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Content.")
        doc.save(str(pdf_path))
        doc.close()

        parser = MultiFormatParser()
        result = await parser.parse(str(pdf_path))

        assert result.metadata.get("content_type") == ".pdf"
        assert result.metadata.get("source_path") == str(pdf_path)


class TestDOCXParserBlocks:
    """DOCX parser returns ParsedDocument with structured blocks."""

    async def test_docx_returns_parsed_document(self, tmp_path):
        """DOCX parser returns ParsedDocument."""
        from src.adapters.parsers import MultiFormatParser

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Simple paragraph.")
        doc.save(str(docx_path))

        parser = MultiFormatParser()
        result = await parser.parse(str(docx_path))

        assert isinstance(result, ParsedDocument)
        assert len(result.blocks) > 0

    async def test_docx_headings_detected(self, tmp_path):
        """DOCX parser produces Heading blocks for heading styles."""
        from src.adapters.parsers import MultiFormatParser

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "headings.docx"
        doc = Document()
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("Body text.")
        doc.add_heading("Section 1.1", level=2)
        doc.add_paragraph("More text.")
        doc.save(str(docx_path))

        parser = MultiFormatParser()
        result = await parser.parse(str(docx_path))

        headings = [b for b in result.blocks if isinstance(b, Heading)]
        assert len(headings) >= 2
        assert headings[0].text == "Chapter One"
        assert headings[0].level == 1
        assert headings[1].text == "Section 1.1"
        assert headings[1].level == 2

    async def test_docx_tables_detected(self, tmp_path):
        """DOCX parser produces TableBlock for tables."""
        from src.adapters.parsers import MultiFormatParser

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "tables.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Age"
        table.cell(1, 0).text = "Alice"
        table.cell(1, 1).text = "30"
        doc.save(str(docx_path))

        parser = MultiFormatParser()
        result = await parser.parse(str(docx_path))

        tables = [b for b in result.blocks if isinstance(b, TableBlock)]
        assert len(tables) >= 1
        assert tables[0].headers == ["Name", "Age"]
        assert tables[0].rows == [["Alice", "30"]]


class TestHTMLParserBlocks:
    """HTML parser returns ParsedDocument with structured blocks."""

    async def test_html_returns_parsed_document(self, tmp_path):
        """HTML parser returns ParsedDocument."""
        from src.adapters.parsers import MultiFormatParser

        html_path = tmp_path / "test.html"
        html_path.write_text(
            "<html><body><p>Simple paragraph.</p></body></html>",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(html_path))

        assert isinstance(result, ParsedDocument)
        assert len(result.blocks) > 0

    async def test_html_headings_detected(self, tmp_path):
        """HTML parser produces Heading blocks for h1-h6 tags."""
        from src.adapters.parsers import MultiFormatParser

        html_path = tmp_path / "headings.html"
        html_path.write_text(
            "<html><body>"
            "<h1>Main Title</h1>"
            "<p>Intro text.</p>"
            "<h2>Section</h2>"
            "<p>Section text.</p>"
            "</body></html>",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(html_path))

        headings = [b for b in result.blocks if isinstance(b, Heading)]
        assert len(headings) >= 2
        assert headings[0].text == "Main Title"
        assert headings[0].level == 1
        assert headings[1].text == "Section"
        assert headings[1].level == 2

    async def test_html_lists_detected(self, tmp_path):
        """HTML parser produces ListBlock for ul/ol tags."""
        from src.adapters.parsers import MultiFormatParser

        html_path = tmp_path / "lists.html"
        html_path.write_text(
            "<html><body>"
            "<ul><li>Item 1</li><li>Item 2</li></ul>"
            "<ol><li>First</li><li>Second</li></ol>"
            "</body></html>",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(html_path))

        lists = [b for b in result.blocks if isinstance(b, ListBlock)]
        assert len(lists) >= 2
        # Unordered list
        assert len(lists[0].items) == 2
        assert lists[0].items[0].text == "Item 1"
        assert lists[0].items[0].ordered is False
        # Ordered list
        assert lists[1].items[0].ordered is True

    async def test_html_tables_detected(self, tmp_path):
        """HTML parser produces TableBlock for table tags."""
        from src.adapters.parsers import MultiFormatParser

        html_path = tmp_path / "tables.html"
        html_path.write_text(
            "<html><body>"
            "<table><thead><tr><th>Name</th><th>Age</th></tr></thead>"
            "<tbody><tr><td>Alice</td><td>30</td></tr></tbody></table>"
            "</body></html>",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(html_path))

        tables = [b for b in result.blocks if isinstance(b, TableBlock)]
        assert len(tables) >= 1
        assert tables[0].headers == ["Name", "Age"]
        assert tables[0].rows == [["Alice", "30"]]


class TestPlainParserBlocks:
    """TXT/MD parser returns ParsedDocument with structured blocks."""

    async def test_plain_returns_parsed_document(self, tmp_path):
        """Plain text parser returns ParsedDocument."""
        from src.adapters.parsers import MultiFormatParser

        txt_path = tmp_path / "test.txt"
        txt_path.write_text("Simple text content.", encoding="utf-8")

        parser = MultiFormatParser()
        result = await parser.parse(str(txt_path))

        assert isinstance(result, ParsedDocument)
        assert len(result.blocks) > 0

    async def test_markdown_headings_detected(self, tmp_path):
        """MD parser produces Heading blocks for # prefix."""
        from src.adapters.parsers import MultiFormatParser

        md_path = tmp_path / "test.md"
        md_path.write_text(
            "# Title\n\n"
            "Intro paragraph.\n\n"
            "## Section\n\n"
            "Section content.\n\n"
            "### Subsection\n\n"
            "Sub content.\n",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(md_path))

        headings = [b for b in result.blocks if isinstance(b, Heading)]
        assert len(headings) >= 3
        assert headings[0].level == 1
        assert headings[0].text == "Title"
        assert headings[1].level == 2
        assert headings[1].text == "Section"
        assert headings[2].level == 3
        assert headings[2].text == "Subsection"

    async def test_markdown_paragraphs_detected(self, tmp_path):
        """MD parser produces Paragraph blocks for text."""
        from src.adapters.parsers import MultiFormatParser

        md_path = tmp_path / "paras.md"
        md_path.write_text(
            "# Title\n\n"
            "First paragraph.\n\n"
            "Second paragraph.\n",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(md_path))

        paragraphs = [b for b in result.blocks if isinstance(b, Paragraph)]
        assert len(paragraphs) >= 2

    async def test_txt_no_headings(self, tmp_path):
        """TXT parser produces only Paragraph blocks (no headings)."""
        from src.adapters.parsers import MultiFormatParser

        txt_path = tmp_path / "plain.txt"
        txt_path.write_text(
            "First paragraph.\n\n"
            "Second paragraph.\n",
            encoding="utf-8",
        )

        parser = MultiFormatParser()
        result = await parser.parse(str(txt_path))

        headings = [b for b in result.blocks if isinstance(b, Heading)]
        assert len(headings) == 0
        paragraphs = [b for b in result.blocks if isinstance(b, Paragraph)]
        assert len(paragraphs) >= 1
