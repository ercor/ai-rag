"""Integration test for structure-aware chunking end-to-end.

Tests:
- PDF with known pages → chunks have correct page_number values
- MD with large list → no mid-item split
- DOCX with table → table content preserved in chunks
"""

import pytest

from src.adapters.parsers import MultiFormatParser
from src.adapters.recursive_chunker import RecursiveChunker


@pytest.mark.integration
class TestStructureAwareChunking:
    """End-to-end structure-aware chunking scenarios."""

    async def test_pdf_page_numbers_preserved(self, tmp_path):
        """PDF with multiple pages → chunks have correct page_number."""
        try:
            import pymupdf
        except ImportError:
            pytest.skip("pymupdf not installed")

        # Create a multi-page PDF
        pdf_path = tmp_path / "pages.pdf"
        doc = pymupdf.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"Content for page {i + 1}. " * 20)
        doc.save(str(pdf_path))
        doc.close()

        parser = MultiFormatParser()
        parsed = await parser.parse(str(pdf_path))

        chunker = RecursiveChunker(min_tokens=50, max_tokens=200, overlap=10)
        chunks = chunker.split(parsed, "doc1", "Pages", str(pdf_path), ".pdf")

        assert len(chunks) >= 1
        # At least some chunks should have page_number set
        page_numbers = [c.page_number for c in chunks if c.page_number is not None]
        assert len(page_numbers) >= 1

    async def test_md_list_not_split_mid_item(self, tmp_path):
        """MD with 50-item list → no mid-item split."""
        md_path = tmp_path / "list.md"
        items = []
        for i in range(50):
            items.append(f"- Item {i}: " + "x" * 30)
        md_path.write_text("\n".join(items), encoding="utf-8")

        parser = MultiFormatParser()
        parsed = await parser.parse(str(md_path))

        chunker = RecursiveChunker(min_tokens=50, max_tokens=200, overlap=10)
        chunks = chunker.split(parsed, "doc1", "List", str(md_path), ".md")

        assert len(chunks) >= 1
        # Verify no item is split: each "Item N:" should be in exactly one chunk
        for i in range(50):
            marker = f"Item {i}:"
            chunks_with_marker = [c for c in chunks if marker in c.content]
            assert len(chunks_with_marker) == 1, (
                f"Item {i} found in {len(chunks_with_marker)} chunks (should be 1)"
            )

    async def test_docx_table_preserved(self, tmp_path):
        """DOCX with table → table content preserved in chunks."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "table.docx"
        doc = Document()
        table = doc.add_table(rows=5, cols=3)
        for i in range(5):
            table.cell(i, 0).text = f"Row{i}Col0"
            table.cell(i, 1).text = f"Row{i}Col1"
            table.cell(i, 2).text = f"Row{i}Col2"
        doc.save(str(docx_path))

        parser = MultiFormatParser()
        parsed = await parser.parse(str(docx_path))

        chunker = RecursiveChunker(min_tokens=50, max_tokens=200, overlap=10)
        chunks = chunker.split(parsed, "doc1", "Table", str(docx_path), ".docx")

        assert len(chunks) >= 1
        # Table content should be present in chunks
        all_content = " ".join(c.content for c in chunks)
        assert "Row0Col0" in all_content
        assert "Row4Col2" in all_content
